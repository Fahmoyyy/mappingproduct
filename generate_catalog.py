"""
Generate catalog.json dari PLN_HOME_Master_Aset.xlsx (sheet 'Jenis Aset' + 'Kelompok Aset').
Jalankan ulang script ini kalau ada produk baru di master aset -- lihat brief di
BRIEF_Web_Mapping_Produk_RAB.md untuk aturan lengkap.

Usage: python generate_catalog.py [path_ke_master_aset.xlsx] [path_output_catalog.json]
"""
import json
import re
import sys
import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\aga\Downloads\PLN_HOME_Master_Aset.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "catalog.json"

UNIT_MAP = {
    "box": "Box", "service": "Service", "stel": "Stel", "paket": "Paket",
    "mandays": "Mandays", "meter": "m", "buah": "Buah", "kali": "Kali",
    "psc": "Pcs", "set": "Set", "1": "Units", "kilogram": "kg", "m": "m",
    "ls": "Ls", "unit": "Units", "orang": "Orang", "pcs": "Pcs", "kg": "kg",
    "rim": "Rim", "batang": "Batang", "liter": "L", "au": "Units",
    "pohon": "Pohon", "bh": "Buah", "lot": "Lot", "kms": "km", "dus": "Dus",
    "botol": "Botol", "pack": "Pack", "psg": "Pasang", "hari": "Hari", "u": "Units",
}

CODE_SUFFIX_RE = re.compile(r"\s*\([^)]*\)\s*$")


def normalize_unit(raw):
    if raw is None:
        return ""
    key = str(raw).strip().lower()
    if key in UNIT_MAP:
        return UNIT_MAP[key]
    return str(raw).strip().title()


def clean_name(name):
    return CODE_SUFFIX_RE.sub("", name).strip()


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    kelompok_ws = wb["Kelompok Aset"]
    kelompok_map = {}
    for row in kelompok_ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        kelompok_map[row[0]] = row[1]

    jenis_ws = wb["Jenis Aset"]
    header = next(jenis_ws.iter_rows(min_row=1, max_row=1, values_only=True))
    col = {name: idx for idx, name in enumerate(header)}

    catalog = []
    for row in jenis_ws.iter_rows(min_row=2, values_only=True):
        if row[col["ACTIVE"]] != 1:
            continue
        name = row[col["KETERANGAN"]]
        if not name:
            continue
        ref = row[col["NORMALISASI"]]
        ref = "" if ref is None else str(ref)
        kelompok_id = row[col["KELOMPOK_ASET"]]
        item = {
            "name": str(name).strip(),
            "nameClean": clean_name(str(name)),
            "ref": ref,
            "unit": normalize_unit(row[col["SATUAN"]]),
            "kelompok": kelompok_map.get(kelompok_id, "") or "",
            "kategori": row[col["KET_KATEGORI_ASET"]] or "",
        }
        catalog.append(item)

    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(catalog, f, ensure_ascii=False, indent=None, separators=(",", ":"))

    print(f"Wrote {len(catalog)} items to {OUT}")


if __name__ == "__main__":
    main()
